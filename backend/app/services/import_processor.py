"""
Import processor service for SlabHub.

This module handles:
1. Image processing from SlabCrop watch folder (duplicate detection, archiving)
2. Metadata import from CSV and XLSX files (Phase 2)

Features:
- CSV/XLSX file parsing with header validation
- Row-level validation with clear error reporting
- Upsert logic (insert or update based on SlabID)
- Import logging and statistics
"""

import csv
import io
import logging
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any, Tuple

from sqlalchemy.orm import Session

from backend.app.config import settings
from backend.app.models import Slab, ImportLog
from backend.app.utils import calculate_perceptual_hash, is_duplicate

logger = logging.getLogger(__name__)


# ============================================================================
# Column Mapping Configuration (Phase 2)
# ============================================================================

# Required columns for import validation
REQUIRED_COLUMNS = ["name"]

# Column mappings: CSV Header (case-insensitive) -> Slab model field
COLUMN_MAPPING = {
    # Identifiers
    "slabid": "public_id",
    "slab_id": "public_id",
    "public_id": "public_id",
    "id": "public_id",

    # Required fields
    "name": "name",
    "title": "name",
    "slab_name": "name",

    # Stone properties
    "stonetype": "stone_type",
    "stone_type": "stone_type",
    "type": "stone_type",
    "material": "stone_type",

    "supplier": "supplier",
    "vendor": "supplier",

    "finish": "finish",
    "surface": "finish",

    "thickness": "thickness",
    "thick": "thickness",

    "color": "color",
    "colour": "color",

    # Dimensions
    "length": "length",
    "len": "length",

    "width": "width",
    "w": "width",

    "squarefeet": "square_feet",
    "square_feet": "square_feet",
    "sqft": "square_feet",
    "area": "square_feet",

    # Status and location
    "location": "location",
    "warehouse": "location",
    "storage": "location",

    "status": "status",
    "state": "status",

    "quantity": "quantity",
    "qty": "quantity",

    # Metadata
    "tags": "tags",
    "categories": "tags",

    "notes": "notes",
    "comments": "notes",
    "description": "notes",

    # Pricing
    "cost": "cost",
    "costprice": "cost",
    "cost_price": "cost",

    "price": "price",
    "sellingprice": "price",
    "selling_price": "price",
}

# Numeric fields for type validation
NUMERIC_FIELDS = {"thickness", "length", "width", "square_feet", "cost", "price"}
INTEGER_FIELDS = {"quantity"}


class ImportProcessor:
    """
    Processes imported images for the slab inventory system.

    Handles image validation, duplicate detection, database record creation,
    file archiving, and import logging. Designed to process images from the
    SlabCrop watch folder.
    """

    def __init__(self, db: Session, batch_id: Optional[str] = None):
        """
        Initialize the import processor.

        Args:
            db: SQLAlchemy database session
            batch_id: Optional batch ID for grouping imports.
                     If not provided, a new UUID-based batch ID is generated.
        """
        self.db = db
        self.batch_id = batch_id or self._generate_batch_id()

        # Initialize counters
        self.processed_count = 0
        self.success_count = 0
        self.failed_count = 0
        self.skipped_count = 0
        self.errors: list[dict] = []

        # Create import log entry
        self.import_log = ImportLog(
            batch_id=self.batch_id,
            import_type="slabcrop_watch",
            source_path=str(settings.slabcrop_output_folder),
            status="running",
            items_processed=0,
            items_success=0,
            items_failed=0,
            items_skipped=0,
            errors=[],
            warnings=[]
        )
        self.db.add(self.import_log)
        self.db.commit()

        logger.info(f"ImportProcessor initialized with batch_id: {self.batch_id}")

    def _generate_batch_id(self) -> str:
        """
        Generate a unique batch ID with timestamp.

        Returns:
            Batch ID string in format: slabcrop_YYYYMMDD_HHMMSS_UUID
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        unique_id = str(uuid.uuid4())[:8]
        return f"slabcrop_{timestamp}_{unique_id}"

    def process_image(self, image_path: Path) -> bool:
        """
        Process a single image file.

        Calculates perceptual hash, checks for duplicates, creates database
        record, and archives the file.

        Args:
            image_path: Path to the image file to process

        Returns:
            True if processing succeeded, False if failed or skipped
        """
        self.processed_count += 1

        try:
            logger.info(f"Processing image: {image_path.name}")

            # Calculate perceptual hash
            phash = calculate_perceptual_hash(image_path)
            if phash is None:
                error_msg = f"Failed to calculate perceptual hash: {image_path.name}"
                logger.error(error_msg)
                self._log_error(image_path, error_msg)
                self.failed_count += 1
                return False

            logger.debug(f"Calculated perceptual hash for {image_path.name}: {phash}")

            # Check for duplicates
            duplicate_slab = self._find_duplicate(phash)
            if duplicate_slab:
                logger.warning(
                    f"Duplicate image detected: {image_path.name} "
                    f"(matches slab {duplicate_slab.public_id})"
                )
                # Archive to duplicates folder
                self._archive_file(image_path, "duplicates")
                self.skipped_count += 1
                return False

            # Extract slab name from filename
            slab_name = self._extract_slab_name(image_path)

            # Generate unique public ID
            public_id = self._generate_public_id()

            # Archive file to processed folder (before creating DB record)
            archived_path = self._archive_file(image_path, "processed")
            if archived_path is None:
                error_msg = f"Failed to archive file: {image_path.name}"
                logger.error(error_msg)
                self._log_error(image_path, error_msg)
                self.failed_count += 1
                return False

            # Create Slab record
            slab = Slab(
                public_id=public_id,
                name=slab_name,
                primary_image=str(archived_path),
                perceptual_hash=phash,
                import_source="slabcrop_watch",
                import_batch_id=self.batch_id,
                status="available",
                quantity=1
            )

            self.db.add(slab)
            self.db.commit()

            logger.info(
                f"Successfully created slab record: {public_id} ({slab_name})"
            )
            self.success_count += 1
            return True

        except Exception as e:
            error_msg = f"Error processing {image_path.name}: {str(e)}"
            logger.error(error_msg, exc_info=True)
            self._log_error(image_path, error_msg)
            self.failed_count += 1

            # Rollback transaction on error
            self.db.rollback()
            return False

        finally:
            # Update import log statistics
            self._update_import_log()

    def _find_duplicate(self, phash: str) -> Optional[Slab]:
        """
        Find duplicate slab by perceptual hash.

        Args:
            phash: Perceptual hash to search for

        Returns:
            Slab record if duplicate found, None otherwise
        """
        try:
            # Get all slabs with perceptual hashes
            slabs_with_hashes = self.db.query(Slab).filter(
                Slab.perceptual_hash.isnot(None)
            ).all()

            # Compare hashes using configured threshold
            threshold = settings.duplicate_threshold

            for slab in slabs_with_hashes:
                if is_duplicate(phash, slab.perceptual_hash, threshold=threshold):
                    logger.debug(
                        f"Found duplicate: {slab.public_id} (name: {slab.name})"
                    )
                    return slab

            return None

        except Exception as e:
            logger.error(f"Error finding duplicates: {e}", exc_info=True)
            return None

    def _extract_slab_name(self, image_path: Path) -> str:
        """
        Extract slab name from image filename.

        Removes file extension and cleans up the filename.

        Args:
            image_path: Path to the image file

        Returns:
            Cleaned slab name string
        """
        # Remove file extension
        name = image_path.stem

        # Replace underscores and hyphens with spaces
        name = name.replace('_', ' ').replace('-', ' ')

        # Clean up multiple spaces
        name = ' '.join(name.split())

        # Capitalize first letter of each word
        name = name.title()

        return name or "Unnamed Slab"

    def _archive_file(self, file_path: Path, subdir: str) -> Optional[Path]:
        """
        Archive file to the archive folder with date-based organization.

        Archives files to: archive_folder/subdir/YYYYMMDD/filename
        Handles filename conflicts by appending a counter.

        Args:
            file_path: Path to the file to archive
            subdir: Subdirectory name (e.g., 'processed' or 'duplicates')

        Returns:
            Path to archived file if successful, None if failed
        """
        try:
            # Build archive path with date subdirectory
            archive_base = Path(settings.processed_archive_folder)
            date_str = datetime.now().strftime("%Y%m%d")
            archive_dir = archive_base / subdir / date_str

            # Create archive directory
            archive_dir.mkdir(parents=True, exist_ok=True)

            # Build destination path
            dest_path = archive_dir / file_path.name

            # Handle filename conflicts
            if dest_path.exists():
                counter = 1
                stem = file_path.stem
                suffix = file_path.suffix

                while dest_path.exists():
                    new_name = f"{stem}_{counter}{suffix}"
                    dest_path = archive_dir / new_name
                    counter += 1
                    if counter > 1000:  # Prevent infinite loops
                        logger.error(
                            f"Too many filename conflicts for {file_path.name}"
                        )
                        return None

                logger.debug(
                    f"Renamed {file_path.name} to {dest_path.name} "
                    f"to avoid conflict"
                )

            # Copy file to archive
            shutil.copy2(file_path, dest_path)
            logger.debug(f"Archived {file_path.name} to {dest_path}")

            # Delete original file
            file_path.unlink()
            logger.debug(f"Deleted original file: {file_path.name}")

            return dest_path

        except Exception as e:
            logger.error(
                f"Error archiving file {file_path.name}: {e}",
                exc_info=True
            )
            return None

    def _log_error(self, file_path: Path, error: str) -> None:
        """
        Log an error for a specific file.

        Args:
            file_path: Path to the file that caused the error
            error: Error message
        """
        error_entry = {
            "file": file_path.name,
            "error": error,
            "timestamp": datetime.now().isoformat()
        }
        self.errors.append(error_entry)
        logger.debug(f"Logged error for {file_path.name}: {error}")

    def _update_import_log(self) -> None:
        """
        Update import log statistics in the database.

        Updates the ImportLog record with current processing statistics.
        """
        try:
            self.import_log.items_processed = self.processed_count
            self.import_log.items_success = self.success_count
            self.import_log.items_failed = self.failed_count
            self.import_log.items_skipped = self.skipped_count
            self.import_log.errors = self.errors if self.errors else None

            self.db.commit()
            logger.debug("Updated import log statistics")

        except Exception as e:
            logger.error(f"Error updating import log: {e}", exc_info=True)
            self.db.rollback()

    def _generate_public_id(self) -> str:
        """
        Generate a unique public ID for a slab.

        Returns:
            Unique public ID string (8-character UUID)
        """
        while True:
            public_id = str(uuid.uuid4())[:8].upper()

            # Check if ID already exists
            existing = self.db.query(Slab).filter(
                Slab.public_id == public_id
            ).first()

            if not existing:
                return public_id

    def process_metadata_csv(self, csv_content: str) -> Dict[str, Any]:
        """
        Process a CSV file containing slab metadata.

        Args:
            csv_content: String content of the CSV file

        Returns:
            Dictionary with results: imported, updated, errors
        """
        self.import_log.import_type = "csv_metadata"
        self.db.commit()

        results = {"imported": 0, "updated": 0, "errors": []}
        
        try:
            # Use StringIO to treat string as a file
            f = io.StringIO(csv_content)
            
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []

            # Required columns validation (case-insensitive check)
            required_columns = ["Name", "Location"]
            header_map = {h.lower(): h for h in headers}
            
            missing = [col for col in required_columns if col.lower() not in header_map]
            if missing:
                error_msg = f"Missing required columns: {', '.join(missing)}"
                self.import_log.status = "failed"
                self.import_log.errors = [{"row": 0, "column": None, "message": error_msg}]
                self.db.commit()
                return {"imported": 0, "updated": 0, "errors": self.import_log.errors}

            # Field mapping (CSV Column -> Slab Model Field)
            # Keys are CSV column names (expected), values are model fields
            mapping = {
                "SlabID": "public_id",
                "Name": "name",
                "StoneType": "stone_type",
                "Supplier": "supplier",
                "Finish": "finish",
                "Thickness": "thickness",
                "Color": "color",
                "Length": "length",
                "Width": "width",
                "SquareFeet": "square_feet",
                "Location": "location",
                "Status": "status",
                "Quantity": "quantity",
                "Tags": "tags",
                "Notes": "notes",
                "Cost": "cost",
                "Price": "price"
            }
            
            # Create a case-insensitive reverse mapping for the reader
            csv_to_model = {}
            for csv_col, model_field in mapping.items():
                if csv_col.lower() in header_map:
                    csv_to_model[header_map[csv_col.lower()]] = model_field

            for i, row in enumerate(reader, start=2):
                self.processed_count += 1
                row_errors = []
                
                # Validation using mapped headers
                name_col = header_map.get("name")
                loc_col = header_map.get("location")
                
                if not row.get(name_col):
                    row_errors.append({"row": i, "column": "Name", "message": "Name is required"})
                if not row.get(loc_col):
                    row_errors.append({"row": i, "column": "Location", "message": "Location is required"})

                # Type validation for numeric fields
                numeric_fields = ["thickness", "length", "width", "square_feet", "cost", "price"]
                # We need to find which CSV columns map to these numeric fields
                model_to_csv = {v: k for k, v in csv_to_model.items()}
                
                for field in numeric_fields:
                    csv_col = model_to_csv.get(field)
                    if csv_col:
                        val = row.get(csv_col)
                        if val:
                            try:
                                float(str(val).replace(',', ''))
                            except ValueError:
                                row_errors.append({"row": i, "column": csv_col, "message": f"Invalid numeric value: {val}"})

                qty_col = model_to_csv.get("quantity")
                if qty_col:
                    val = row.get(qty_col)
                    if val:
                        try:
                            int(val)
                        except ValueError:
                            row_errors.append({"row": i, "column": qty_col, "message": f"Invalid integer: {val}"})

                if row_errors:
                    results["errors"].extend(row_errors)
                    self.errors.extend(row_errors)
                    self.failed_count += 1
                    continue

                # Upsert Logic
                try:
                    pub_id_col = model_to_csv.get("public_id")
                    public_id = row.get(pub_id_col) if pub_id_col else None
                    
                    existing_slab = None
                    if public_id:
                        existing_slab = self.db.query(Slab).filter(Slab.public_id == public_id).first()

                    slab_data = {}
                    extra_json = {}

                    for col_name, value in row.items():
                        if value is None or str(value).strip() == "":
                            continue
                            
                        model_field = csv_to_model.get(col_name)
                        if model_field:
                            # Convert types
                            if model_field in numeric_fields:
                                slab_data[model_field] = float(str(value).replace(',', ''))
                            elif model_field == "quantity":
                                slab_data[model_field] = int(value)
                            else:
                                slab_data[model_field] = value
                        else:
                            extra_json[col_name] = value

                    if existing_slab:
                        # Update
                        for field, value in slab_data.items():
                            setattr(existing_slab, field, value)
                        if extra_json:
                            existing_slab.extra_json = extra_json
                        results["updated"] += 1
                        self.success_count += 1
                    else:
                        # Insert
                        if not slab_data.get("public_id"):
                            slab_data["public_id"] = self._generate_public_id()
                        
                        new_slab = Slab(
                            **slab_data,
                            extra_json=extra_json if extra_json else None,
                            import_source="csv_metadata",
                            import_batch_id=self.batch_id
                        )
                        self.db.add(new_slab)
                        results["imported"] += 1
                        self.success_count += 1

                    # Periodic commit for large imports
                    if (results["imported"] + results["updated"]) % 50 == 0:
                        self.db.commit()

                except Exception as e:
                    err = {"row": i, "column": None, "message": str(e)}
                    results["errors"].append(err)
                    self.errors.append(err)
                    self.failed_count += 1
                    self.db.rollback()

            self.db.commit()
            self._update_import_log()
            return results
            
        except Exception as e:
            logger.error(f"Critical error processing CSV: {e}", exc_info=True)
            self.import_log.status = "failed"
            self.import_log.errors = [{"row": 0, "column": None, "message": f"Critical error: {str(e)}"}]
            self.db.commit()
            return {"imported": 0, "updated": 0, "errors": self.import_log.errors}

    def process_metadata_file(self) -> bool:
        """
        Process metadata from the configured file path in settings.

        Returns:
            True if successful, False otherwise
        """
        if not settings.import_metadata_file:
            logger.error("import_metadata_file setting is not configured")
            return False
            
        file_path = Path(settings.import_metadata_file)
        if not file_path.exists():
            logger.error(f"Metadata file not found: {file_path}")
            return False
            
        try:
            logger.info(f"Processing metadata file: {file_path}")
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
                results = self.process_metadata_csv(content)
                logger.info(f"File import results: {results['imported']} imported, {results['updated']} updated, {len(results['errors'])} errors")
                return True
        except Exception as e:
            logger.error(f"Error reading metadata file {file_path}: {e}", exc_info=True)
            return False

    def finalize(self) -> None:
        """
        Finalize the import process and complete the import log.

        Updates the import log status to 'completed' or 'failed' based on
        processing results and adds a summary.
        """
        try:
            # Determine final status
            if self.failed_count > 0 and self.success_count == 0:
                status = "failed"
            elif self.failed_count > 0:
                status = "partial"
            else:
                status = "completed"

            # Generate summary
            summary = (
                f"SlabCrop import completed. "
                f"Processed: {self.processed_count}, "
                f"Success: {self.success_count}, "
                f"Failed: {self.failed_count}, "
                f"Skipped (duplicates): {self.skipped_count}"
            )

            # Update import log
            self.import_log.status = status
            self.import_log.completed_at = datetime.now()
            self.import_log.summary = summary

            self.db.commit()

            logger.info(f"Import finalized: {summary}")

        except Exception as e:
            logger.error(f"Error finalizing import: {e}", exc_info=True)
            self.db.rollback()

    # ========================================================================
    # Phase 2: Enhanced Metadata Import (CSV/XLSX)
    # ========================================================================

    def process_file_upload(
        self,
        file_content: bytes,
        filename: str
    ) -> Dict[str, Any]:
        """
        Process an uploaded CSV or XLSX file with enhanced validation.

        Args:
            file_content: Raw bytes of the uploaded file
            filename: Original filename (used to detect file type)

        Returns:
            Dict with import results including:
            - batch_id, status, file_type
            - rows_processed, rows_imported, rows_updated, rows_failed, rows_skipped
            - errors (list of {row, column, message})
            - warnings (list of {row, column, message})
            - summary (human readable)
        """
        file_ext = Path(filename).suffix.lower()

        if file_ext == ".csv":
            return self._process_csv_enhanced(file_content)
        elif file_ext in (".xlsx", ".xls"):
            return self._process_xlsx(file_content)
        else:
            return {
                "batch_id": self.batch_id,
                "status": "failed",
                "file_type": file_ext,
                "rows_processed": 0,
                "rows_imported": 0,
                "rows_updated": 0,
                "rows_failed": 0,
                "rows_skipped": 0,
                "errors": [{"row": 0, "column": None, "message": f"Unsupported file type: {file_ext}. Use .csv or .xlsx"}],
                "warnings": [],
                "summary": f"Import failed: unsupported file type {file_ext}"
            }

    def _process_csv_enhanced(self, file_content: bytes) -> Dict[str, Any]:
        """
        Process CSV file with enhanced header mapping and validation.

        Args:
            file_content: Raw bytes of CSV file

        Returns:
            Import results dictionary
        """
        self.import_log.import_type = "csv_metadata"
        self.db.commit()

        results = self._init_results("csv")

        try:
            # Decode CSV content
            try:
                text = file_content.decode("utf-8")
            except UnicodeDecodeError:
                text = file_content.decode("latin-1", errors="replace")

            f = io.StringIO(text)
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []

            # Validate and map headers
            header_mapping, header_errors = self._validate_headers(headers)
            if header_errors:
                results["errors"].extend(header_errors)
                results["status"] = "failed"
                results["summary"] = "Import failed: header validation errors"
                self._finalize_import_log(results)
                return results

            # Process rows
            self._process_rows(reader, header_mapping, results)

            # Finalize
            self._finalize_import_log(results)
            return results

        except Exception as e:
            logger.error(f"Critical error processing CSV: {e}", exc_info=True)
            results["status"] = "failed"
            results["errors"].append({"row": 0, "column": None, "message": f"Critical error: {str(e)}"})
            results["summary"] = f"Import failed: {str(e)}"
            self._finalize_import_log(results)
            return results

    def _process_xlsx(self, file_content: bytes) -> Dict[str, Any]:
        """
        Process XLSX file with header mapping and validation.

        Args:
            file_content: Raw bytes of XLSX file

        Returns:
            Import results dictionary
        """
        self.import_log.import_type = "xlsx_metadata"
        self.db.commit()

        results = self._init_results("xlsx")

        try:
            # Import openpyxl (optional dependency)
            try:
                import openpyxl
            except ImportError:
                results["status"] = "failed"
                results["errors"].append({
                    "row": 0,
                    "column": None,
                    "message": "XLSX support requires openpyxl package. Install with: pip install openpyxl"
                })
                results["summary"] = "Import failed: openpyxl not installed"
                self._finalize_import_log(results)
                return results

            # Load workbook from bytes
            wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
            ws = wb.active

            if ws is None:
                results["status"] = "failed"
                results["errors"].append({"row": 0, "column": None, "message": "No active worksheet found in XLSX file"})
                self._finalize_import_log(results)
                return results

            # Get headers from first row
            headers = []
            for cell in ws[1]:
                val = cell.value
                if val is not None:
                    headers.append(str(val).strip())
                else:
                    headers.append("")

            # Validate and map headers
            header_mapping, header_errors = self._validate_headers(headers)
            if header_errors:
                results["errors"].extend(header_errors)
                results["status"] = "failed"
                results["summary"] = "Import failed: header validation errors"
                self._finalize_import_log(results)
                return results

            # Convert XLSX rows to dict format (like CSV reader)
            def xlsx_row_generator():
                for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    row_dict = {}
                    for col_idx, value in enumerate(row):
                        if col_idx < len(headers) and headers[col_idx]:
                            row_dict[headers[col_idx]] = value
                    yield row_dict

            # Process rows
            self._process_rows(xlsx_row_generator(), header_mapping, results)

            wb.close()

            # Finalize
            self._finalize_import_log(results)
            return results

        except Exception as e:
            logger.error(f"Critical error processing XLSX: {e}", exc_info=True)
            results["status"] = "failed"
            results["errors"].append({"row": 0, "column": None, "message": f"Critical error: {str(e)}"})
            results["summary"] = f"Import failed: {str(e)}"
            self._finalize_import_log(results)
            return results

    def _init_results(self, file_type: str) -> Dict[str, Any]:
        """Initialize results dictionary."""
        return {
            "batch_id": self.batch_id,
            "status": "completed",
            "file_type": file_type,
            "rows_processed": 0,
            "rows_imported": 0,
            "rows_updated": 0,
            "rows_failed": 0,
            "rows_skipped": 0,
            "errors": [],
            "warnings": [],
            "summary": ""
        }

    def _validate_headers(self, headers: List[str]) -> Tuple[Dict[str, str], List[Dict]]:
        """
        Validate and map CSV/XLSX headers to model fields.

        Args:
            headers: List of header names from the file

        Returns:
            Tuple of (header_mapping, errors)
            - header_mapping: Dict mapping original header -> model field
            - errors: List of header validation errors
        """
        errors = []
        header_mapping = {}  # original_header -> model_field

        if not headers:
            errors.append({"row": 1, "column": None, "message": "No headers found in file"})
            return header_mapping, errors

        # Map headers to model fields
        for original_header in headers:
            if not original_header or not original_header.strip():
                continue

            normalized = original_header.lower().strip().replace(" ", "_").replace("-", "_")
            model_field = COLUMN_MAPPING.get(normalized)

            if model_field:
                header_mapping[original_header] = model_field
                logger.debug(f"Mapped header '{original_header}' -> '{model_field}'")

        # Check required columns
        mapped_fields = set(header_mapping.values())
        for required in REQUIRED_COLUMNS:
            if required not in mapped_fields:
                errors.append({
                    "row": 1,
                    "column": None,
                    "message": f"Missing required column: '{required}'. "
                               f"Accepted headers: {', '.join(k for k, v in COLUMN_MAPPING.items() if v == required)}"
                })

        if not header_mapping:
            errors.append({
                "row": 1,
                "column": None,
                "message": "No recognized columns found. Check column headers against expected format."
            })

        return header_mapping, errors

    def _process_rows(
        self,
        row_iterator,
        header_mapping: Dict[str, str],
        results: Dict[str, Any]
    ) -> None:
        """
        Process rows from CSV or XLSX with validation and upsert logic.

        Args:
            row_iterator: Iterator of row dictionaries
            header_mapping: Dict mapping original header -> model field
            results: Results dictionary to update in place
        """
        # Create reverse mapping: model_field -> original_header
        model_to_header = {v: k for k, v in header_mapping.items()}

        for row_num, row in enumerate(row_iterator, start=2):
            results["rows_processed"] += 1
            self.processed_count += 1

            # Skip completely empty rows
            if not any(v for v in row.values() if v is not None and str(v).strip()):
                results["rows_skipped"] += 1
                self.skipped_count += 1
                continue

            # Validate row
            row_errors = self._validate_row(row, header_mapping, row_num)
            if row_errors:
                results["errors"].extend(row_errors)
                results["rows_failed"] += 1
                self.failed_count += 1
                continue

            # Process row (upsert)
            try:
                imported, updated = self._upsert_slab(row, header_mapping, row_num, results)
                if imported:
                    results["rows_imported"] += 1
                    self.success_count += 1
                elif updated:
                    results["rows_updated"] += 1
                    self.success_count += 1

                # Periodic commit
                if self.success_count % 50 == 0:
                    self.db.commit()

            except Exception as e:
                logger.error(f"Error processing row {row_num}: {e}", exc_info=True)
                results["errors"].append({"row": row_num, "column": None, "message": str(e)})
                results["rows_failed"] += 1
                self.failed_count += 1
                self.db.rollback()

        # Final commit
        self.db.commit()

    def _validate_row(
        self,
        row: Dict[str, Any],
        header_mapping: Dict[str, str],
        row_num: int
    ) -> List[Dict]:
        """
        Validate a single row's data.

        Args:
            row: Row data dictionary (original headers as keys)
            header_mapping: Dict mapping original header -> model field
            row_num: Row number for error reporting

        Returns:
            List of validation errors (empty if valid)
        """
        errors = []
        model_to_header = {v: k for k, v in header_mapping.items()}

        # Check required fields
        for required in REQUIRED_COLUMNS:
            header = model_to_header.get(required)
            if header:
                value = row.get(header)
                if value is None or str(value).strip() == "":
                    errors.append({
                        "row": row_num,
                        "column": header,
                        "message": f"Required field '{required}' is empty"
                    })

        # Validate numeric fields
        for field in NUMERIC_FIELDS:
            header = model_to_header.get(field)
            if header:
                value = row.get(header)
                if value is not None and str(value).strip():
                    try:
                        float(str(value).replace(",", "").strip())
                    except ValueError:
                        errors.append({
                            "row": row_num,
                            "column": header,
                            "message": f"Invalid number: '{value}'"
                        })

        # Validate integer fields
        for field in INTEGER_FIELDS:
            header = model_to_header.get(field)
            if header:
                value = row.get(header)
                if value is not None and str(value).strip():
                    try:
                        int(float(str(value).replace(",", "").strip()))
                    except ValueError:
                        errors.append({
                            "row": row_num,
                            "column": header,
                            "message": f"Invalid integer: '{value}'"
                        })

        return errors

    def _upsert_slab(
        self,
        row: Dict[str, Any],
        header_mapping: Dict[str, str],
        row_num: int,
        results: Dict[str, Any]
    ) -> Tuple[bool, bool]:
        """
        Insert or update a slab from row data.

        Args:
            row: Row data dictionary
            header_mapping: Dict mapping original header -> model field
            row_num: Row number for warnings
            results: Results dict for adding warnings

        Returns:
            Tuple of (was_imported, was_updated)
        """
        model_to_header = {v: k for k, v in header_mapping.items()}
        slab_data = {}
        extra_json = {}

        # Extract data from row
        for original_header, value in row.items():
            if value is None or str(value).strip() == "":
                continue

            model_field = header_mapping.get(original_header)
            if model_field:
                # Convert types
                if model_field in NUMERIC_FIELDS:
                    slab_data[model_field] = float(str(value).replace(",", "").strip())
                elif model_field in INTEGER_FIELDS:
                    slab_data[model_field] = int(float(str(value).replace(",", "").strip()))
                else:
                    slab_data[model_field] = str(value).strip()
            else:
                # Store unmapped columns in extra_json
                if original_header.strip():
                    extra_json[original_header] = value

        # Check for existing slab by public_id
        public_id = slab_data.get("public_id")
        existing_slab = None

        if public_id:
            existing_slab = self.db.query(Slab).filter(Slab.public_id == public_id).first()

        if existing_slab:
            # Update existing slab
            for field, value in slab_data.items():
                if field != "public_id":  # Don't update public_id
                    setattr(existing_slab, field, value)

            if extra_json:
                # Merge extra_json
                current_extra = existing_slab.extra_json or {}
                current_extra.update(extra_json)
                existing_slab.extra_json = current_extra

            logger.debug(f"Updated slab {public_id} from row {row_num}")
            return False, True

        else:
            # Create new slab
            if not slab_data.get("public_id"):
                slab_data["public_id"] = self._generate_public_id()

            # Set defaults
            if "status" not in slab_data:
                slab_data["status"] = "available"
            if "quantity" not in slab_data:
                slab_data["quantity"] = 1

            new_slab = Slab(
                **slab_data,
                extra_json=extra_json if extra_json else None,
                import_source=f"{results['file_type']}_metadata",
                import_batch_id=self.batch_id
            )
            self.db.add(new_slab)

            logger.debug(f"Created slab {slab_data['public_id']} from row {row_num}")
            return True, False

    def _finalize_import_log(self, results: Dict[str, Any]) -> None:
        """
        Finalize the import log with results.

        Args:
            results: Import results dictionary
        """
        try:
            # Determine status
            if results["rows_failed"] > 0 and results["rows_imported"] == 0 and results["rows_updated"] == 0:
                results["status"] = "failed"
            elif results["rows_failed"] > 0:
                results["status"] = "partial"
            else:
                results["status"] = "completed"

            # Generate summary
            results["summary"] = (
                f"Import {results['status']}: "
                f"{results['rows_processed']} processed, "
                f"{results['rows_imported']} imported, "
                f"{results['rows_updated']} updated, "
                f"{results['rows_failed']} failed, "
                f"{results['rows_skipped']} skipped"
            )

            # Update import log
            self.import_log.status = results["status"]
            self.import_log.items_processed = results["rows_processed"]
            self.import_log.items_success = results["rows_imported"] + results["rows_updated"]
            self.import_log.items_failed = results["rows_failed"]
            self.import_log.items_skipped = results["rows_skipped"]
            self.import_log.errors = results["errors"] if results["errors"] else None
            self.import_log.warnings = results["warnings"] if results["warnings"] else None
            self.import_log.summary = results["summary"]
            self.import_log.completed_at = datetime.now()

            self.db.commit()

            logger.info(f"Import finalized: {results['summary']}")

        except Exception as e:
            logger.error(f"Error finalizing import log: {e}", exc_info=True)
            self.db.rollback()
