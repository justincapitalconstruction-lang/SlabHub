#!/usr/bin/env python3
"""
Test script for Phase 2 Metadata Import functionality.

Tests:
1. CSV import with valid data
2. CSV import with validation errors
3. Header validation
4. XLSX import (if openpyxl available)

Usage:
    python scripts/test_metadata_import.py
"""

import io
import sys
from pathlib import Path

# Ensure project root is in path
project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))


def create_test_csv_valid():
    """Create a valid CSV for testing."""
    return """Name,StoneType,Supplier,Thickness,Location,Status,Quantity,Price
Test Marble 001,Marble,Stone Supplier Inc,2.0,Warehouse A,available,1,150.00
Test Granite 002,Granite,Rock Quarry LLC,3.0,Warehouse B,available,2,250.00
Test Quartzite 003,Quartzite,Premier Stone,2.5,Warehouse A,reserved,1,350.00
"""


def create_test_csv_with_errors():
    """Create a CSV with validation errors for testing."""
    return """Name,StoneType,Thickness,Location,Quantity,Price
,Marble,2.0,Warehouse A,1,150.00
Test Slab 002,Granite,invalid_number,Warehouse B,2,250.00
Test Slab 003,Quartzite,2.5,Warehouse C,not_an_int,350.00
"""


def create_test_csv_missing_headers():
    """Create a CSV with no recognized headers."""
    return """Foo,Bar,Baz
value1,value2,value3
"""


def test_csv_import():
    """Test CSV import functionality."""
    from backend.app.models import SessionLocal, Slab
    from backend.app.services.import_processor import ImportProcessor

    print("\n[TEST] CSV Import with Valid Data")
    print("-" * 40)

    db = SessionLocal()
    try:
        # Create processor
        processor = ImportProcessor(db)

        # Process valid CSV
        csv_content = create_test_csv_valid()
        results = processor.process_file_upload(csv_content.encode("utf-8"), "test.csv")

        print(f"  Batch ID: {results['batch_id']}")
        print(f"  Status: {results['status']}")
        print(f"  Rows Processed: {results['rows_processed']}")
        print(f"  Rows Imported: {results['rows_imported']}")
        print(f"  Rows Updated: {results['rows_updated']}")
        print(f"  Rows Failed: {results['rows_failed']}")
        print(f"  Errors: {len(results['errors'])}")
        print(f"  Summary: {results['summary']}")

        # Verify import
        if results["rows_imported"] == 3 and results["rows_failed"] == 0:
            print("  [OK] Valid CSV import successful")

            # Clean up test data
            db.query(Slab).filter(Slab.import_batch_id == results["batch_id"]).delete()
            db.commit()
            return True
        else:
            print("  [FAIL] Expected 3 imports, 0 failures")
            return False

    except Exception as e:
        print(f"  [FAIL] Exception: {e}")
        return False
    finally:
        db.close()


def test_csv_validation_errors():
    """Test CSV import with validation errors."""
    from backend.app.models import SessionLocal, Slab
    from backend.app.services.import_processor import ImportProcessor

    print("\n[TEST] CSV Import with Validation Errors")
    print("-" * 40)

    db = SessionLocal()
    try:
        processor = ImportProcessor(db)

        csv_content = create_test_csv_with_errors()
        results = processor.process_file_upload(csv_content.encode("utf-8"), "test_errors.csv")

        print(f"  Status: {results['status']}")
        print(f"  Rows Processed: {results['rows_processed']}")
        print(f"  Rows Failed: {results['rows_failed']}")
        print(f"  Errors: {len(results['errors'])}")

        for err in results["errors"]:
            print(f"    Row {err['row']}, Column {err['column']}: {err['message']}")

        # We expect:
        # Row 2: Name is empty (required field)
        # Row 3: Thickness is invalid number
        # Row 4: Quantity is invalid integer
        if results["rows_failed"] == 3 and len(results["errors"]) >= 3:
            print("  [OK] Validation errors detected correctly")

            # Clean up any test data that got through
            db.query(Slab).filter(Slab.import_batch_id == results["batch_id"]).delete()
            db.commit()
            return True
        else:
            print(f"  [FAIL] Expected 3 failures with 3+ errors")
            return False

    except Exception as e:
        print(f"  [FAIL] Exception: {e}")
        return False
    finally:
        db.close()


def test_header_validation():
    """Test header validation for missing required columns."""
    from backend.app.models import SessionLocal
    from backend.app.services.import_processor import ImportProcessor

    print("\n[TEST] Header Validation - Missing Required Columns")
    print("-" * 40)

    db = SessionLocal()
    try:
        processor = ImportProcessor(db)

        csv_content = create_test_csv_missing_headers()
        results = processor.process_file_upload(csv_content.encode("utf-8"), "test_headers.csv")

        print(f"  Status: {results['status']}")
        print(f"  Errors: {len(results['errors'])}")

        for err in results["errors"]:
            print(f"    Row {err['row']}: {err['message']}")

        if results["status"] == "failed" and any("name" in str(err["message"]).lower() for err in results["errors"]):
            print("  [OK] Missing required column detected")
            return True
        else:
            print("  [FAIL] Should have failed with missing 'name' column")
            return False

    except Exception as e:
        print(f"  [FAIL] Exception: {e}")
        return False
    finally:
        db.close()


def test_xlsx_import():
    """Test XLSX import (if openpyxl available)."""
    print("\n[TEST] XLSX Import Support")
    print("-" * 40)

    try:
        import openpyxl
        print("  openpyxl is installed")

        from backend.app.models import SessionLocal, Slab
        from backend.app.services.import_processor import ImportProcessor

        # Create a simple XLSX in memory
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Slabs"

        # Add headers
        headers = ["Name", "StoneType", "Supplier", "Thickness", "Location", "Status"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        # Add data row
        data = ["XLSX Test Slab", "Marble", "Test Supplier", 2.5, "Warehouse X", "available"]
        for col, value in enumerate(data, 1):
            ws.cell(row=2, column=col, value=value)

        # Save to bytes
        xlsx_buffer = io.BytesIO()
        wb.save(xlsx_buffer)
        xlsx_content = xlsx_buffer.getvalue()

        db = SessionLocal()
        try:
            processor = ImportProcessor(db)
            results = processor.process_file_upload(xlsx_content, "test.xlsx")

            print(f"  Status: {results['status']}")
            print(f"  File Type: {results['file_type']}")
            print(f"  Rows Imported: {results['rows_imported']}")
            print(f"  Errors: {len(results['errors'])}")

            if results["rows_imported"] == 1:
                print("  [OK] XLSX import successful")

                # Clean up
                db.query(Slab).filter(Slab.import_batch_id == results["batch_id"]).delete()
                db.commit()
                return True
            else:
                print("  [FAIL] Expected 1 import")
                return False

        except Exception as e:
            print(f"  [FAIL] Exception: {e}")
            return False
        finally:
            db.close()

    except ImportError:
        print("  [SKIP] openpyxl not installed - XLSX support unavailable")
        return None


def test_import_result_schema():
    """Test that ImportResult schema is properly defined."""
    print("\n[TEST] ImportResult Schema")
    print("-" * 40)

    try:
        from backend.app.schemas import ImportResult, ImportRowError

        # Create a test result
        result = ImportResult(
            batch_id="test_batch_123",
            status="completed",
            file_type="csv",
            rows_processed=10,
            rows_imported=8,
            rows_updated=2,
            rows_failed=0,
            rows_skipped=0,
            errors=[],
            warnings=[],
            summary="Test import completed"
        )

        print(f"  batch_id: {result.batch_id}")
        print(f"  status: {result.status}")
        print("  [OK] ImportResult schema works correctly")
        return True

    except Exception as e:
        print(f"  [FAIL] Exception: {e}")
        return False


def main():
    """Run all import tests."""
    print("=" * 50)
    print("SlabHub Phase 2 - Metadata Import Tests")
    print("=" * 50)

    results = []

    # Run tests
    results.append(("ImportResult Schema", test_import_result_schema()))
    results.append(("Header Validation", test_header_validation()))
    results.append(("CSV Validation Errors", test_csv_validation_errors()))
    results.append(("CSV Valid Import", test_csv_import()))
    results.append(("XLSX Import", test_xlsx_import()))

    # Summary
    print("\n" + "=" * 50)
    print("Test Summary")
    print("=" * 50)

    passed = 0
    failed = 0
    skipped = 0

    for name, result in results:
        if result is True:
            print(f"  [PASS] {name}")
            passed += 1
        elif result is False:
            print(f"  [FAIL] {name}")
            failed += 1
        else:
            print(f"  [SKIP] {name}")
            skipped += 1

    print()
    print(f"Passed: {passed}, Failed: {failed}, Skipped: {skipped}")

    if failed > 0:
        sys.exit(1)
    else:
        sys.exit(0)


if __name__ == "__main__":
    main()
